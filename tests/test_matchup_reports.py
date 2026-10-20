import hashlib
import json
import math
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.special import expit
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import StandardScaler

import UFC_ML_Model as core
from matchup_reports import explain_winner, render_card_markdown, _votes, _feature_group
from scripts.predict_card_reports import validate_model_metadata


class SignalModel:
    def __init__(self, sign=1):
        self.sign = sign

    def predict_proba(self, x):
        z = self.sign * (2 * x["d_sig_str_pm"] - x["d_td_p15"])
        p = expit(z + .5 * x["is_title"])
        return np.column_stack([1 - p, p])


class SquareCalibrator:
    def predict_proba(self, x):
        p = .1 + .8 * np.asarray(x)[:, 0] ** 2
        return np.column_stack([1 - p, p])


def make_ensemble():
    cols = ["d_sig_str_pm", "d_td_p15", "elo_r", "elo_b", "is_title"]
    background = pd.DataFrame([[0, 0, 1500, 1500, 0], [1, 1, 1600, 1400, 1]], columns=cols)
    imputer = SimpleImputer().fit(background)
    scaler = StandardScaler().fit(background)
    return core.SuperEnsembleModel(
        {"signal": SignalModel(), "unused": SignalModel(-1)}, imputer, scaler, cols,
        {"kind": "weighted", "model_order": ["signal", "unused"],
         "weights": {"signal": 1.0, "unused": 0.0}},
    )


class AttributionTests(unittest.TestCase):
    def test_distance_metrics_are_not_mislabeled_as_stance(self):
        self.assertEqual(_feature_group("d_distance_pct"), "striking")
        self.assertEqual(_feature_group("d_rd1_distance"), "cardio")
        self.assertEqual(_feature_group("same_stance"), "physical")

    def setUp(self):
        self.ensemble = make_ensemble()
        self.row = pd.DataFrame([dict(d_sig_str_pm=1., d_td_p15=.4,
                                     elo_r=1500., elo_b=1500., is_title=0.)])

    def test_exact_two_group_shapley_against_closed_form(self):
        report, _ = explain_winner(self.ensemble, self.row, core)
        factors = {f["group"]: f["contribution_probability_a"] for f in report["factors"]}
        strike = .5 * ((expit(2) - .5) + (expit(1.6) - expit(-.4)))
        wrestling = .5 * ((expit(-.4) - .5) + (expit(1.6) - expit(2)))
        self.assertAlmostEqual(factors["striking"], strike, places=12)
        self.assertAlmostEqual(factors["wrestling"], wrestling, places=12)
        self.assertAlmostEqual(report["baseline_probability_a"], .5, places=12)
        self.assertAlmostEqual(report["model_probability_a"], expit(1.6), places=12)
        self.assertEqual(report["evaluated_coalitions"], 4)
        self.assertAlmostEqual(report["reconstruction_residual"], 0, places=12)

    def test_corner_swap_reverses_contributions_and_preserves_pairs(self):
        self.row["elo_r"], self.row["elo_b"] = 1700., 1400.
        a, _ = explain_winner(self.ensemble, self.row, core)
        b, _ = explain_winner(self.ensemble, core._swap_features(self.row), core)
        self.assertAlmostEqual(a["model_probability_a"] + b["model_probability_a"], 1, places=12)
        by_group = {f["group"]: f for f in b["factors"]}
        for f in a["factors"]:
            self.assertAlmostEqual(f["contribution_probability_a"],
                                   -by_group[f["group"]]["contribution_probability_a"], places=12)
        rating = next(f for f in a["factors"] if f["group"] == "ratings")
        self.assertEqual(set(rating["feature_names"]), {"elo_r", "elo_b"})

    def test_calibrated_baseline_is_measured_not_assumed_half(self):
        self.ensemble.calibrator = SquareCalibrator()
        report, _ = explain_winner(self.ensemble, self.row, core)
        self.assertAlmostEqual(report["baseline_probability_a"], .3, places=12)
        self.assertAlmostEqual(report["model_probability_a"], .1+.8*expit(1.6)**2, places=12)
        actual = self.ensemble.predict_proba_single(self.row.iloc[0].to_dict())
        self.assertAlmostEqual(report["model_probability_a"], actual, places=12)

    def test_missing_values_and_corner_correction_use_actual_inference_chain(self):
        self.row.loc[0, "d_sig_str_pm"] = np.nan
        self.ensemble.corner_correction = .15
        report, _ = explain_winner(self.ensemble, self.row, core)
        actual = self.ensemble.predict_proba_single(self.row.iloc[0].to_dict())
        self.assertAlmostEqual(report["model_probability_a"], actual, places=12)
        self.assertTrue(math.isfinite(report["baseline_probability_a"]))
        self.assertAlmostEqual(report["reconstruction_residual"], 0, places=12)

    def test_zero_weight_model_is_excluded_from_disagreement(self):
        _, probabilities = explain_winner(self.ensemble, self.row, core)
        votes, disagreement = _votes(self.ensemble, probabilities, "A", "B")
        self.assertFalse(votes[1]["active"])
        self.assertEqual(disagreement["active_models"], 1)
        self.assertEqual(disagreement["weighted_standard_deviation"], 0)

    def test_nonlinear_combiner_uses_actual_output_and_descriptive_vote_weights(self):
        class Stacker:
            def predict_proba(self, x):
                p = expit(np.asarray(x)[:, 0] * 2 - .5)
                return np.column_stack([1-p, p])
        self.ensemble.combiner = {"kind":"stacker", "model":Stacker(),
                                  "model_order":["signal","unused"]}
        report, probabilities = explain_winner(self.ensemble, self.row, core)
        self.assertAlmostEqual(report["model_probability_a"],
            self.ensemble.predict_proba_single(self.row.iloc[0].to_dict()), places=12)
        votes, disagreement = _votes(self.ensemble, probabilities, "A", "B")
        self.assertTrue(all(v["ensemble_weight"] is None for v in votes))
        self.assertIn("nonlinear", disagreement["weighting"])


class IntegrationTests(unittest.TestCase):
    def setUp(self):
        self.pipeline = core.UFCSuperModelPipeline()
        self.pipeline.model = make_ensemble()
        for name, age in (("Fighter Alpha", 30), ("Fighter Beta", 35)):
            self.pipeline.fighter_history = self.pipeline.fighter_history or {}
            self.pipeline.fighter_history[name] = [core._make_synthetic_fight_record(
                pd.Timestamp("2025-01-01"), {"age": age})]
        self.pipeline.glicko_ratings = {}
        self.pipeline.opp_glicko_list = {}

    def predict(self, **kwargs):
        return self.pipeline.predict_matchup("Fighter Alpha", "Fighter Beta",
            weight_class="Welterweight", gender="Men", rounds=5,
            prediction_date="2026-10-18", **kwargs)

    def test_report_does_not_change_prediction_and_is_strict_json(self):
        plain = self.predict()
        explained = self.predict(include_explanation=True)
        report = explained.pop("explanation")
        self.assertEqual(plain, explained)
        self.assertEqual(report["coverage"][0]["recorded_fights"], 1)
        self.assertTrue(report["coverage"][0]["flags"])
        json.dumps(report, allow_nan=False)

    def test_title_context_is_passed_explicitly(self):
        report = self.predict(include_explanation=True, is_title=True)["explanation"]
        self.assertTrue(report["bout_context"]["is_title"])
        # Directly validate the model's title-sensitive forward/reverse chain.
        row = pd.DataFrame([dict(d_sig_str_pm=1., d_td_p15=.4,
                                elo_r=1500., elo_b=1500., is_title=1.)])
        attribution, _ = explain_winner(self.pipeline.model, row, core)
        self.assertAlmostEqual(attribution["model_probability_a"],
            (expit(2-.4+.5)+1-expit(-2+.4+.5))/2, places=12)

    def test_past_prediction_cannot_use_future_history(self):
        with self.assertRaisesRegex(ValueError, "historical model snapshot"):
            self.pipeline.predict_matchup("Fighter Alpha", "Fighter Beta", prediction_date="2024-01-01")

    def test_debutant_coverage_is_explicit(self):
        p = self.pipeline.predict_matchup("New Fighter", "Fighter Beta", include_explanation=True,
                                          prediction_date="2026-10-18")
        coverage = p["explanation"]["coverage"][0]
        self.assertEqual(coverage["recorded_fights"], 0)
        self.assertIn("generic debutant priors", coverage["flags"][0])

    def test_workbook_preserves_predictions_and_adds_report_sheet(self):
        plain = self.predict()
        explained = self.predict(include_explanation=True)
        with tempfile.TemporaryDirectory() as directory:
            before, after = Path(directory)/"before.xlsx", Path(directory)/"after.xlsx"
            core.export_to_excel(before, [plain], [])
            core.export_to_excel(after, [explained], [])
            a,b = load_workbook(before), load_workbook(after)
            self.assertEqual(list(a["Predictions"].values), list(b["Predictions"].values))
            self.assertEqual(b["Explanations"].max_row, 2)
            self.assertIn("Not a confidence interval", b["Explanations"]["D2"].value)
            a.close();b.close()
        text = render_card_markdown({"event":"Example", "prediction_date":"2026-10-18",
            "data_cutoff":"2026-10-17", "bouts":[{"fighters":["A","B"],"prediction":explained}]})
        self.assertIn("Model factor group", text)

    def test_odds_blending_does_not_get_attributed_to_fighter_statistics(self):
        p = self.predict(include_explanation=True, red_odds=-500, blue_odds=400)
        report = p["explanation"]
        self.assertTrue(report["market"]["odds_blended"])
        self.assertAlmostEqual(report["winner_attribution"]["model_probability_a"], p["prob_a_model"])
        self.assertNotAlmostEqual(p["prob_a"], p["prob_a_model"])
        self.assertAlmostEqual(report["market"]["expected_return_a"], p["prob_a_model"] * 1.2 - 1)
        self.assertAlmostEqual(report["market"]["expected_return_b"], p["prob_b_model"] * 5 - 1)

    def test_no_odds_produces_no_invented_market_edges(self):
        market = self.predict(include_explanation=True)["explanation"]["market"]
        self.assertIsNone(market["expected_return_a"])
        self.assertIsNone(market["expected_return_b"])
        self.assertIsNone(market["model_edge_vs_no_vig_a"])

    def test_corrected_gender_context_and_inactivity_do_not_raise_old_flags(self):
        p = self.pipeline.predict_matchup("Fighter Alpha", "Fighter Beta",
            weight_class="Bantamweight", gender="Women", prediction_date="2026-10-18",
            include_explanation=True)
        report = p["explanation"]
        self.assertFalse(any("Gender context" in flag for flag in report["flags"]))
        coverage = report["coverage"][0]
        self.assertEqual(coverage["days_since_last_fight"], 655)
        self.assertFalse(any("recorded dates imply" in flag for flag in coverage["flags"]))
        inactivity = next(row for row in report["comparison"] if row["feature"] == "days_inactive")
        self.assertEqual(inactivity["a"],655)


class ProvenanceTests(unittest.TestCase):
    def test_source_or_data_mismatch_is_rejected_before_loading(self):
        source = b"SETTING = 1\n"
        metadata = {"source_data_sha256": hashlib.sha256(source+b"data").hexdigest(),
                    "ufc_environment":{}}
        with patch.dict(os.environ, {}, clear=True):
            validate_model_metadata(metadata, source, b"data", source)
            with self.assertRaisesRegex(ValueError, "metadata"):
                validate_model_metadata(metadata, source, b"changed data", source)
            with self.assertRaisesRegex(ValueError, "Training code"):
                validate_model_metadata(metadata, source, b"data", b"SETTING = 2\n")
        with patch.dict(os.environ, {"UFC_MOV_ENABLED":"0"}, clear=True):
            with self.assertRaisesRegex(ValueError, "configuration"):
                validate_model_metadata(metadata, source, b"data", source)


if __name__ == "__main__":
    unittest.main()
